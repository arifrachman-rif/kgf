import os
import openpyxl

downloads = "/mnt/c/Users/rifra/Downloads"

cashout_files = [f for f in os.listdir(downloads) if "cash out" in f.lower() and f.endswith(".xlsx")]
print(f"Found {len(cashout_files)} Cash Out excel files:")

for fname in cashout_files:
    fpath = os.path.join(downloads, fname)
    print(f"\n==========================================")
    print(f"FILE: {fname}")
    print(f"==========================================")
    try:
        wb = openpyxl.load_workbook(fpath, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            print(f"\n--- Sheet: {sname} ---")
            for r in ws.iter_rows(values_only=True):
                non_empty = [str(cell).strip() for cell in r if cell is not None and str(cell).strip() != ""]
                if non_empty:
                    print("  ", " | ".join(non_empty))
    except Exception as e:
        print(f"Error: {e}")
