import os
import openpyxl
import docx

downloads = "/mnt/c/Users/rifra/Downloads"

print("--- EXCEL FILES CHECK ---")
excel_files = [f for f in os.listdir(downloads) if f.startswith("Mutiara Cash Out") and f.endswith(".xlsx")]
for ef in excel_files:
    path = os.path.join(downloads, ef)
    print(f"\nReading {ef}:")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    line = " | ".join(vals)
                    if any(word in line.lower() for word in ['nama', 'enum', 'rp', 'bank', 'transfer', 'mutiara']):
                        print("  ", line[:100])
    except Exception as e:
        print(" Error reading excel:", e)

print("\n--- DOCX FILES CHECK ---")
docx_files = [f for f in os.listdir(downloads) if "enumerator" in f.lower() or "mutiara" in f.lower() if f.endswith(".docx")]
for df in docx_files:
    path = os.path.join(downloads, df)
    print(f"\nReading {df}:")
    try:
        doc = docx.Document(path)
        for p in doc.paragraphs:
            if p.text.strip():
                print("  ", p.text[:100])
        for t in doc.tables:
            for r in t.rows:
                row_txt = [c.text.strip() for c in r.cells]
                print("   Table row:", " | ".join(row_txt))
    except Exception as e:
        print(" Error reading docx:", e)
