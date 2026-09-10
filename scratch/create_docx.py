import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_paragraph_with_spacing(doc, text="", style=None, before=0, after=6):
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p

def create_styled_docx():
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styling
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33) # Off-black

    # Color Palette
    PRIMARY_COLOR = RGBColor(0x00, 0x33, 0x66) # Dark Blue
    SECONDARY_COLOR = RGBColor(0x70, 0x80, 0x90) # Slate Grey
    FILL_HEADER_HEX = "003366"
    FILL_ROW_HEX = "F0F4F8"

    # Paths translated for WSL
    md_path = "/mnt/c/Users/rifra/Downloads/robusta_lampung_data_sheet.md"
    dest_path = "/mnt/c/Users/rifra/Downloads/robusta_lampung_data_sheet.docx"
    
    if not os.path.exists(md_path):
        # fallback to relative path inside workspace
        md_path = "data/robusta_lampung_data_sheet.md"
        if not os.path.exists(md_path):
            md_path = "C:/Users/rifra/Downloads/robusta_lampung_data_sheet.md"
            dest_path = "C:/Users/rifra/Downloads/robusta_lampung_data_sheet.docx"

    print(f"Reading from: {md_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_rows = []
    
    for line in lines:
        line_stripped = line.strip()
        
        # Check if table row
        if line_stripped.startswith('|'):
            in_table = True
            # Parse cells
            cells = [c.strip() for c in line_stripped.split('|')[1:-1]]
            # Skip separator line (e.g., |:---|:---|)
            if all(c.startswith(':') or c.startswith('-') or c == '' for c in cells):
                continue
            table_rows.append(cells)
            continue
        else:
            # If we were in a table, process and write it now
            if in_table and table_rows:
                write_table(doc, table_rows, FILL_HEADER_HEX, FILL_ROW_HEX)
                table_rows = []
                in_table = False
            
            # Skip empty lines
            if not line_stripped:
                continue
                
            # Process markdown markers
            if line_stripped.startswith('# '):
                text = line_stripped[2:]
                h = doc.add_heading(level=1)
                run = h.add_run(text)
                run.font.name = 'Calibri'
                run.font.size = Pt(20)
                run.font.bold = True
                run.font.color.rgb = PRIMARY_COLOR
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(12)
                
            elif line_stripped.startswith('## '):
                text = line_stripped[3:]
                h = doc.add_heading(level=2)
                run = h.add_run(text)
                run.font.name = 'Calibri'
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = PRIMARY_COLOR
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(6)
                
            elif line_stripped.startswith('### '):
                text = line_stripped[4:]
                h = doc.add_heading(level=3)
                run = h.add_run(text)
                run.font.name = 'Calibri'
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = SECONDARY_COLOR
                h.paragraph_format.space_before = Pt(6)
                h.paragraph_format.space_after = Pt(4)
                
            elif line_stripped.startswith('* ') or line_stripped.startswith('- '):
                text = line_stripped[2:]
                p = add_paragraph_with_spacing(doc, style='List Bullet', after=4)
                parse_and_add_inline_formatting(p, text)
                
            elif line_stripped.startswith('>') or line_stripped.startswith('> '):
                text = line_stripped.lstrip('>').strip()
                p = add_paragraph_with_spacing(doc, after=6)
                p.paragraph_format.left_indent = Inches(0.5)
                run = p.add_run(text)
                run.font.italic = True
                run.font.color.rgb = SECONDARY_COLOR
                
            elif line_stripped.startswith('---'):
                # Add horizontal line equivalent (border or simple spacer)
                add_paragraph_with_spacing(doc, after=12)
                
            else:
                p = add_paragraph_with_spacing(doc, after=6)
                parse_and_add_inline_formatting(p, line_stripped)

    # If document ends with a table
    if in_table and table_rows:
        write_table(doc, table_rows, FILL_HEADER_HEX, FILL_ROW_HEX)

    doc.save(dest_path)
    print(f"File saved successfully to {dest_path}")

def parse_and_add_inline_formatting(paragraph, text):
    # Basic markdown inline parser for bold **text** and *italic*
    parts = []
    i = 0
    n = len(text)
    
    while i < n:
        if text[i:i+2] == '**':
            # bold
            end = text.find('**', i+2)
            if end != -1:
                parts.append((text[i+2:end], True, False))
                i = end + 2
            else:
                parts.append((text[i:], False, False))
                break
        elif text[i] == '*':
            # italic
            end = text.find('*', i+1)
            if end != -1:
                parts.append((text[i+1:end], False, True))
                i = end + 1
            else:
                parts.append((text[i:], False, False))
                break
        else:
            # plain text up to next formatting marker
            next_bold = text.find('**', i)
            next_italic = text.find('*', i)
            
            # find the nearest
            markers = [m for m in [next_bold, next_italic] if m != -1]
            if not markers:
                parts.append((text[i:], False, False))
                break
            else:
                next_marker = min(markers)
                parts.append((text[i:next_marker], False, False))
                i = next_marker
                
    for content, bold, italic in parts:
        run = paragraph.add_run(content)
        if bold:
            run.bold = True
        if italic:
            run.italic = True

def write_table(doc, rows, header_color, row_color):
    if not rows:
        return
        
    num_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'
    
    # Format table headers
    for c_idx, cell_text in enumerate(rows[0]):
        cell = table.cell(0, c_idx)
        cell.text = cell_text
        set_cell_background(cell, header_color)
        # Font format
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White text
                
    # Format data rows
    for r_idx, row_data in enumerate(rows[1:]):
        r_num = r_idx + 1
        # Strip markers like **text** from table data for clean look
        for c_idx, cell_text in enumerate(row_data):
            cell = table.cell(r_num, c_idx)
            # Remove markdown bold markers inside table cells for cleaner look
            clean_text = cell_text.replace('**', '').replace('*', '')
            cell.text = clean_text
            
            # Shading alternate rows
            if r_num % 2 == 1:
                set_cell_background(cell, row_color)
                
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)

    # Add spacing after table
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

if __name__ == '__main__':
    create_styled_docx()
