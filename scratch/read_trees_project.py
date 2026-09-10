import openpyxl

file_path = "data/Daftar Petani Agroforestry 2026 (1).xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

if 'trees Planting Project' in wb.sheetnames:
    sheet = wb['trees Planting Project']
    print("--- trees Planting Project Sheet Dump ---")
    for r in sheet.iter_rows(values_only=True):
        if any(r):
            print([str(x).strip() if x is not None else "" for x in r][:12])
