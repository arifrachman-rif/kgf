import os
import openpyxl

downloads = "/mnt/c/Users/rifra/Downloads"

excel_files = [f for f in os.listdir(downloads) if f.endswith(".xlsx")]
print(f"Found {len(excel_files)} excel files in Downloads:")

all_names_found = set()

for ef in excel_files:
    path = os.path.join(downloads, ef)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        print(f"\n==========================================")
        print(f"FILE: {ef}")
        print(f"==========================================")
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            print(f" Sheet: {sheetname}")
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    line = " | ".join(vals)
                    print(f"   Row {r_idx:2d}: {line}")
                    # Collect candidate names
                    for cell in vals:
                        if len(cell.split()) >= 2 and not any(char in cell for char in ['0','1','2','3','4','5','6','7','8','9','@','http','.com','/','-']):
                            all_names_found.add(cell)
    except Exception as e:
        print(f"Error reading {ef}: {e}")

print("\n==========================================")
print("ALL CANDIDATE NAMES FROM EXCEL FILES:")
print("==========================================")
for n in sorted(all_names_found):
    print(" -", n)
