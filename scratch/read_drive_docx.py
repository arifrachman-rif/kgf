import os
import sys
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
import io
from googleapiclient.http import MediaIoBaseDownload
import docx

# Set up paths
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SKILLS_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, '..', '.agent', 'skills'))
TOKEN_FILE = os.path.join(SKILLS_DIR, 'personal-drive-connector', 'token.json')
SCOPES = ['https://www.googleapis.com/auth/drive']

def authenticate():
    if not os.path.exists(TOKEN_FILE):
        print(f"Error: token.json not found at {TOKEN_FILE}", file=sys.stderr)
        return None
    return Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)

def download_and_parse(file_id, output_txt_path=None):
    creds = authenticate()
    if not creds:
        return
    
    service = build('drive', 'v3', credentials=creds)
    
    # Get metadata
    meta = service.files().get(fileId=file_id, fields='name, mimeType', supportsAllDrives=True).execute()
    file_name = meta.get('name')
    mime_type = meta.get('mimeType')
    print(f"--- File: {file_name} ({mime_type}) ---")
    
    # Download binary
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        status, done = downloader.next_chunk()
    
    # Process docx
    fh.seek(0)
    doc = docx.Document(fh)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(para.text)
    
    # Process tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            # remove duplicates caused by merged cells
            cleaned_row = []
            for item in row_text:
                if not cleaned_row or cleaned_row[-1] != item:
                    cleaned_row.append(item)
            fullText.append(" | ".join(cleaned_row))
            
    content = "\n".join(fullText)
    
    if output_txt_path:
        with open(output_txt_path, 'w', encoding='utf-8') as f:
            f.write(content)
        print(f"Content saved to {output_txt_path}")
    else:
        print(content[:5000]) # Print first 5000 chars
        if len(content) > 5000:
            print("\n... [TRUNCATED] ...")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python3 read_drive_docx.py FILE_ID [OUTPUT_TXT_PATH]")
        sys.exit(1)
    file_id = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else None
    download_and_parse(file_id, output_path)
