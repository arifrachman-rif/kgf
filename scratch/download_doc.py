import os
import io
import sys
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

try:
    import docx
except ImportError:
    print("Installing python-docx...")
    os.system("pip install python-docx")
    import docx

def main():
    SCRIPT_DIR = os.path.abspath(os.path.join('.agent', 'skills', 'work-drive-connector'))
    TOKEN_FILE = os.path.join(SCRIPT_DIR, 'token.json')
    
    if not os.path.exists(TOKEN_FILE):
        print(f"Token not found at {TOKEN_FILE}")
        return

    creds = Credentials.from_authorized_user_file(TOKEN_FILE, ['https://www.googleapis.com/auth/drive'])
    service = build('drive', 'v3', credentials=creds)

    file_id = "1e0Jdhw2Vw9yr4c4CrUfBBavl1I8pltzb"
    print(f"Downloading file ID: {file_id}")

    request = service.files().get_media(fileId=file_id)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while done is False:
        status, done = downloader.next_chunk()
        print(f"Download {int(status.progress() * 100)}%.")

    with open('draft.docx', 'wb') as f:
        f.write(fh.getvalue())
        
    print("Download complete. Extracting text...")
    
    doc = docx.Document('draft.docx')
    text = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)
            
    # also extract tables
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.replace('\n', ' '))
            text.append(" | ".join(row_data))
    
    with open('draft.txt', 'w', encoding='utf-8') as f:
        f.write('\n'.join(text))
        
    print("Text extracted successfully to draft.txt!")

if __name__ == "__main__":
    main()
