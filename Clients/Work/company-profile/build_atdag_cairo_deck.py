"""KGF profile untuk ATDAG Cairo: 2 halaman profil + 1 halaman katalog produk.

Dibuat 1 September 2026. Angka ditarik dari 00_CORE_CONTENT_BANK_EN.md.
Nutmeg dan cocoa butter TIDAK ada di portfolio terdokumentasi, jadi ditulis
"sourced to order" tanpa spesifikasi karangan. Halal ditulis apa adanya:
belum dipegang, akan diurus sesuai kebutuhan pembeli.

Ubah di sini lalu jalankan ulang; jangan edit PPTX-nya langsung.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import io, os

A = "assets_from_c3"
OUT = "KGF_Company_Profile_ATDAG_CAIRO.pptx"

# ---------- palet ----------
INK      = RGBColor(0x14, 0x1A, 0x16)
GREEN_D  = RGBColor(0x0E, 0x2A, 0x18)
GREEN    = RGBColor(0x1B, 0x5E, 0x3A)
GREEN_L  = RGBColor(0xE6, 0xEE, 0xE8)
GOLD     = RGBColor(0xC6, 0xA3, 0x1A)
CREAM    = RGBColor(0xFA, 0xF8, 0xF3)
GREY     = RGBColor(0x55, 0x5B, 0x56)
GREY_L   = RGBColor(0xD9, 0xD9, 0xD3)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

HEAD = "Georgia"
BODY = "Segoe UI"

W, H = 13.333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


# ---------- helper ----------
def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    return s


def text(slide, x, y, w, h, blocks, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """blocks: list of (string, size, bold, color, font, space_after, line_spacing)"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, b in enumerate(blocks):
        txt, sz, bold, col, fnt, sa, ls = b
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sa)
        p.line_spacing = ls
        r = p.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold
        r.font.color.rgb = col; r.font.name = fnt
    return tb


def B(txt, sz=9.5, bold=False, col=INK, fnt=BODY, sa=4, ls=1.18):
    return (txt, sz, bold, col, fnt, sa, ls)


def logo(slide, path, x, y, h):
    """Tempel logo dengan tinggi tetap, lebar mengikuti rasio setelah trim."""
    im = Image.open(os.path.join(A, path)).convert("RGBA")
    bb = im.getbbox()
    if bb:
        im = im.crop(bb)
    w = h * im.width / im.height
    buf = io.BytesIO(); im.save(buf, "PNG"); buf.seek(0)
    slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(w), Inches(h))
    return w


def photo(slide, path, x, y, w, h, darken=0):
    """Crop-to-fill lalu tempel. darken 0..255 untuk overlay hijau gelap."""
    im = Image.open(path).convert("RGB")
    tr = w / h
    sr = im.width / im.height
    if sr > tr:
        nw = int(im.height * tr)
        im = im.crop(((im.width - nw) // 2, 0, (im.width - nw) // 2 + nw, im.height))
    else:
        nh = int(im.width / tr)
        top = int((im.height - nh) * 0.4)
        im = im.crop((0, top, im.width, top + nh))
    if darken:
        ov = Image.new("RGB", im.size, (14, 42, 24))
        im = Image.blend(im, ov, darken / 255)
    buf = io.BytesIO(); im.save(buf, "JPEG", quality=88); buf.seek(0)
    slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(w), Inches(h))


def section(slide, x, y, w, label):
    """Judul bagian: garis emas pendek + teks kapital."""
    rect(slide, x, y + 0.055, 0.30, 0.035, fill=GOLD)
    text(slide, x + 0.42, y, w - 0.42, 0.24,
         [B(label.upper(), 9.5, True, GREEN, BODY, 0, 1.0)])


def nlines(txt, w_in, fs, bold=False):
    """Perkiraan jumlah baris setelah wrap. Lebar karakter rata-rata Segoe UI
    ~0.50 em (0.53 kalau tebal); 1 pt = 1/72 inci."""
    em = fs / 72.0
    cpl = max(4, int(w_in / (em * (0.53 if bold else 0.50))))
    n = 0
    for para in str(txt).split("\n"):
        words, line = para.split(), ""
        c = 1
        for wd in words:
            trial = wd if not line else line + " " + wd
            if len(trial) <= cpl:
                line = trial
            else:
                c += 1; line = wd
        n += c
    return n


def table(slide, x, y, w, cols, rows, head_h=0.32, pad=0.20,
          fs=8.6, head_fs=8.2, zebra=True, ls=1.15):
    """Tabel manual dengan tinggi baris dihitung dari hasil wrap sebenarnya."""
    cw = [w * c for c in cols]
    rect(slide, x, y, w, head_h, fill=GREEN)
    cx = x
    for i, c in enumerate(rows[0]):
        text(slide, cx + 0.11, y + 0.07, cw[i] - 0.22, head_h,
             [B(c.upper(), head_fs, True, WHITE, BODY, 0, 1.0)])
        cx += cw[i]
    ry = y + head_h
    lh = fs * ls / 72.0
    for ri, row in enumerate(rows[1:]):
        n = max(nlines(c, cw[i] - 0.24, fs, bold=(i == 0)) for i, c in enumerate(row))
        rh = n * lh + pad
        if zebra and ri % 2 == 0:
            rect(slide, x, ry, w, rh, fill=CREAM)
        rect(slide, x, ry + rh - 0.008, w, 0.008, fill=GREY_L)
        cx = x
        for ci, c in enumerate(row):
            bold = ci == 0
            text(slide, cx + 0.11, ry + pad / 2 - 0.015, cw[ci] - 0.22, rh,
                 [B(str(c), fs, bold, INK if bold else GREY, BODY, 0, ls)])
            cx += cw[ci]
        ry += rh
    return ry


def consortium_strip(slide, y, note=True):
    """Bilah putih berisi logo KGF + mitra. Wordmark aGROWforests & GIZ hitam,
    jadi bilah ini WAJIB terang."""
    rect(slide, 0, y, W, H - y, fill=WHITE)
    rect(slide, 0, y, W, 0.030, fill=GOLD)
    cy = y + 0.26
    lw = logo(slide, "logo_kgf.png", 0.62, cy - 0.02, 0.60)
    rect(slide, 0.62 + lw + 0.40, cy + 0.02, 0.012, 0.54, fill=GREY_L)
    if note:
        text(slide, 0.62 + lw + 0.62, cy + 0.16, 3.0, 0.24,
             [B("Lead local implementer, aGROWforests", 7.4, False, GREY, BODY, 0, 1.0)])
    partners = [("logo_agrowforests_tr.png", 0.52), ("logo_ssii.png", 0.40),
                ("logo_giz_tr.png", 0.36)]
    widths = []
    for p, hh in partners:
        im = Image.open(os.path.join(A, p)).convert("RGBA")
        bb = im.getbbox()
        if bb:
            im = im.crop(bb)
        widths.append(hh * im.width / im.height)
    GAP = 0.62
    total = sum(widths) + GAP * (len(partners) - 1)
    px = W - 0.62 - total
    for (p, hh), pw in zip(partners, widths):
        logo(slide, p, px, cy + (0.60 - hh) / 2, hh)
        px += pw + GAP


# =====================================================================
# HALAMAN 1 — profil
# =====================================================================
s1 = prs.slides.add_slide(BLANK)
rect(s1, 0, 0, W, H, fill=WHITE)

STRIP = 6.34            # bilah logo mulai di sini; semua isi harus di atasnya
HERO = 2.58
photo(s1, "assets_from_c3/cover_bg.jpg", 0, 0, W, HERO, darken=120)
rect(s1, 0, HERO - 0.042, W, 0.042, fill=GOLD)

text(s1, 0.75, 0.40, 7.2, 0.3,
     [B("PT KLUMBAYAN GOLD FARM  ·  LAMPUNG, INDONESIA", 9.5, True,
        RGBColor(0xE8, 0xD9, 0x92), BODY, 0, 1.0)])
text(s1, 0.75, 0.76, 8.4, 1.4,
     [B("Traceable spices and cocoa,", 29, True, WHITE, HEAD, 2, 1.05),
      B("plot to container", 29, True, RGBColor(0xE8, 0xD9, 0x92), HEAD, 0, 1.05)])
text(s1, 0.75, 1.86, 8.7, 0.4,
     [B("1,250 agroforestry smallholders  ·  2,000-tonne facility at Natar  ·  "
        "every lot verified in TAPAK", 10.2, False, RGBColor(0xDC, 0xE6, 0xDD), BODY, 0, 1.2)])

text(s1, 9.75, 1.90, 2.9, 0.4,
     [B("Prepared for", 7.8, True, RGBColor(0xC9, 0xB8, 0x78), BODY, 1, 1.0),
      B("ATDAG Cairo  ·  September 2026", 9.2, True, WHITE, BODY, 0, 1.15)],
     align=PP_ALIGN.RIGHT)

# --- kolom kiri ---
LX, LW = 0.75, 6.30
y = HERO + 0.34
section(s1, LX, y, LW, "Who we are")
text(s1, LX, y + 0.34, LW, 1.6, [B(
    "KGF is an Indonesian spice and cocoa company based in Lampung, the country's "
    "principal pepper-producing province. We are not a broker. We buy directly from "
    "1,250 smallholder farmers we organised and trained across three regencies "
    "— Pesawaran, Tanggamus and Lampung Barat — grade and consolidate at our own "
    "2,000-tonne warehouse in Natar, and ship under contract.", 9.2, False, GREY, BODY, 6, 1.26),
    B("The company was founded in 2022 to restore land degraded by former gold mining "
      "and return villagers to professional farming. That is the origin of the name.",
      9.2, False, GREY, BODY, 0, 1.26)])

y2 = y + 1.72
section(s1, LX, y2, LW, "Two businesses that fund each other")
text(s1, LX, y2 + 0.34, LW, 1.7, [B(
    "Tech-enabled trading.  Buying direct removes the layers of collectors that normally "
    "sit between an Indonesian smallholder and an exporter. That is what makes both the "
    "traceability claim and the farmer price premium real.", 9.2, False, GREY, BODY, 6, 1.26),
    B("Regenerative agribusiness.  Our supply base is agroforestry, not monoculture. Crops "
      "grow in a multi-strata system alongside shade and timber trees, on land being restored "
      "rather than cleared — producing the verified deforestation-free supply that premium "
      "buyers increasingly require.", 9.2, False, GREY, BODY, 0, 1.26)])

# --- kolom kanan: fakta ---
RX, RW = 7.72, 4.86
FH = 2.26
rect(s1, RX, HERO + 0.34, RW, FH, fill=GREEN_L)
rect(s1, RX, HERO + 0.34, 0.045, FH, fill=GREEN)
fy = HERO + 0.50
text(s1, RX + 0.34, fy, RW - 0.7, 0.24,
     [B("AT A GLANCE", 8.6, True, GREEN, BODY, 0, 1.0)])
facts = [("2022", "Established, Pesawaran, Lampung"),
         ("1,250", "Smallholder farmers under contract"),
         ("~600 ha", "Under management (1,250 × 0.5 ha assumed)"),
         ("2,000 t", "Warehouse and processing, Gudang Natar"),
         ("17", "Verified field buying stations, Tanggamus"),
         ("15", "Internal staff, incl. 5 field facilitators")]
fy += 0.31
for k, v in facts:
    text(s1, RX + 0.34, fy - 0.012, 1.15, 0.3, [B(k, 11, True, GREEN_D, HEAD, 0, 1.0)])
    text(s1, RX + 1.50, fy + 0.028, RW - 1.84, 0.3, [B(v, 8.0, False, GREY, BODY, 0, 1.1)])
    fy += 0.29

# --- traksi ---
ty = HERO + 0.34 + FH + 0.22
section(s1, RX, ty, RW, "Where we ship today")
text(s1, RX, ty + 0.30, RW, 0.7, [B(
    "Verstegen Spices and Sauces of the Netherlands contracts 80–100 t of black pepper for "
    "2026 and audited our Natar warehouse in August. Europe is our only export market to "
    "date — Egypt is the second we are here to open.", 8.6, False, GREY, BODY, 0, 1.2)])

consortium_strip(s1, STRIP)
text(s1, 0, 7.20, W, 0.22,
     [B("PT Klumbayan Gold Farm  ·  Company Profile  ·  Page 1 of 3",
        7, False, GREY, BODY, 0, 1.0)], align=PP_ALIGN.CENTER)


# =====================================================================
# HALAMAN 2 — kapabilitas
# =====================================================================
s2 = prs.slides.add_slide(BLANK)
rect(s2, 0, 0, W, H, fill=WHITE)
rect(s2, 0, 0, W, 0.92, fill=GREEN_D)
rect(s2, 0, 0.885, W, 0.035, fill=GOLD)
text(s2, 0.75, 0.24, 8.5, 0.5,
     [B("Capability, stated honestly", 20, True, WHITE, HEAD, 0, 1.0)])
text(s2, 8.6, 0.34, 4.0, 0.3,
     [B("PT KLUMBAYAN GOLD FARM", 8.5, True, RGBColor(0xC9, 0xB8, 0x78), BODY, 0, 1.0)],
     align=PP_ALIGN.RIGHT)

# --- kapasitas ---
y = 1.28
section(s2, 0.75, y, 7.4, "Capacity — the constraint is working capital, not plant")
text(s2, 0.75, y + 0.36, 7.4, 0.5, [B(
    "Our facility at Natar is deliberately not running at its limit. Machine hours are "
    "available and our farmer network can supply more than we currently buy, because until "
    "now we concentrated on a single European contract. New buyer volume displaces nothing.",
    9.3, False, GREY, BODY, 0, 1.28)])

tend = table(s2, 0.75, y + 0.96, 7.4, [0.34, 0.38, 0.28],
             [["Line", "Position today", "Facility"],
              ["Black pepper, premium 580 g/l", "80–100 t contracted for 2026; expandable against a firm 2027 order", "Natar"],
              ["Black pepper, standard 500 g/l", "Available now — 200 t per year", "Natar"],
              ["Robusta coffee, green", "300 t/month sustained, headroom to 720 t/month", "Natar + 17 stations"],
              ["Cocoa, bean and processed", "Available; powder and nibs produced in-house", "Natar"]],
             fs=8.4)

# --- TAPAK ---
ty = tend + 0.34
TH = 1.66
rect(s2, 0.75, ty, 7.4, TH, fill=CREAM)
rect(s2, 0.75, ty, 0.045, TH, fill=GOLD)
text(s2, 1.12, ty + 0.20, 6.85, 1.4,
     [B("TAPAK — traceability that can be checked", 11, True, GREEN_D, HEAD, 6, 1.0),
      B("Traceability and Accountability of Professional Agri-KGF. Our own platform, live and "
        "in production. It records farmer identity, plot polygon coordinates, purchase volume "
        "per transaction, digital receipts, and warehouse lot movement.", 9.2, False, GREY, BODY, 6, 1.26),
      B("A buyer asking which farms a container came from receives a plot-level answer. We can "
        "grant dashboard access for independent verification rather than asking a buyer to take "
        "our word for it.", 9.2, False, GREY, BODY, 0, 1.26)])

# --- kolom kanan ---
RX, RW = 8.62, 3.95
photo(s2, "assets_from_c3/ops_quality_control.jpg", RX, 1.28, RW, 1.72)

sy = 3.22
section(s2, RX, sy, RW, "Sustainability")
text(s2, RX, sy + 0.36, RW, 1.5, [B(
    "aGROWforests — a consortium funded by GIZ, the German development agency — covers land "
    "restoration, deforestation prevention and certified supply chains. GIZ funds it and audits "
    "compliance. KGF is the lead local implementer.", 8.8, False, GREY, BODY, 6, 1.24),
    B("Composting runs at our LKC and Taman Sari facilities plus two communal units in Ulu Belu. "
      "More than 10 tonnes of compost has been distributed to farmers to date.",
      8.8, False, GREY, BODY, 0, 1.24)])

cy = 5.02
section(s2, RX, cy, RW, "Compliance — including what we do not hold")
comp = [("Halal (BPJPH)", "Not yet held. Will be obtained to buyer requirement.", True),
        ("Organic", "Not certified. Land in transition; we do not claim it.", True),
        ("EUDR due diligence", "In preparation with Verstegen and Fairfood.", False),
        ("K3 and waste SOP", "In place at Gudang Natar.", False),
        ("BPJS and payroll", "Paid and active.", False)]
ly = cy + 0.34
for k, v, flag in comp:
    rect(s2, RX, ly + 0.05, 0.065, 0.065, fill=GOLD if flag else GREEN)
    tb = s2.shapes.add_textbox(Inches(RX + 0.19), Inches(ly), Inches(RW - 0.19), Inches(0.40))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.line_spacing = 1.15
    r1 = p.add_run(); r1.text = k + " — "
    r1.font.size = Pt(8.2); r1.font.bold = True; r1.font.color.rgb = INK; r1.font.name = BODY
    r2 = p.add_run(); r2.text = v
    r2.font.size = Pt(8.2); r2.font.color.rgb = GREY; r2.font.name = BODY
    ly += 0.335 if nlines(k + " — " + v, RW - 0.43, 8.2) <= 1 else 0.475

text(s2, 0, 7.16, W, 0.22,
     [B("PT Klumbayan Gold Farm  ·  Company Profile  ·  Page 2 of 3",
        7, False, GREY, BODY, 0, 1.0)], align=PP_ALIGN.CENTER)


# =====================================================================
# HALAMAN 3 — katalog produk
# =====================================================================
s3 = prs.slides.add_slide(BLANK)
rect(s3, 0, 0, W, H, fill=WHITE)
rect(s3, 0, 0, W, 0.92, fill=GREEN_D)
rect(s3, 0, 0.885, W, 0.035, fill=GOLD)
text(s3, 0.75, 0.24, 8.5, 0.5,
     [B("Product catalogue", 20, True, WHITE, HEAD, 0, 1.0)])
text(s3, 8.6, 0.34, 4.0, 0.3,
     [B("ORIGIN: LAMPUNG, INDONESIA", 8.5, True, RGBColor(0xC9, 0xB8, 0x78), BODY, 0, 1.0)],
     align=PP_ALIGN.RIGHT)

y = 1.14
end = table(s3, 0.75, y, 11.83, [0.19, 0.23, 0.30, 0.28],
            [["Product", "Grade / form", "Specification", "Availability"],
             ["Black pepper — Premium", "Premium FAQ, export grade",
              "Bulk density min. 580 g/l. Lampung Black Pepper is a protected origin name.",
              "≈100 t/year, contracted to our European buyer. Expandable against a firm 2027 order."],
             ["Black pepper — Standard", "Grinding and blending grade",
              "Bulk density 500 g/l. Sound, clean pepper below European premium density.",
              "Available now — 200 t/year. Does not compete with the premium line."],
             ["Robusta coffee", "Green bean",
              "Dry natural (asalan kering), and sorted to defect 80 / 120.",
              "300 t/month sustained, headroom to 720 t/month. Sourced Tanggamus."],
             ["Cocoa bean", "Fermented and non-fermented",
              "Fermented carries a premium over non-fermented at the farm gate.",
              "Available across the Lampung network."],
             ["Cocoa processed", "Powder and nibs",
              "Downstream processed in-house at Natar, not raw bean only.",
              "Available."],
             ["Cocoa butter", "Pressed",
              "Specification confirmed against enquiry.",
              "Sourced to order — not a stocked line."],
             ["Nutmeg", "Whole dried",
              "Specification confirmed against enquiry.",
              "Sourced to order — not a stocked line."]],
            fs=8.2, head_fs=8.0, pad=0.17)

# --- catatan lini lain + spesifikasi ---
ny = end + 0.22
NH = 1.06
rect(s3, 0.75, ny, 5.72, NH, fill=CREAM)
rect(s3, 0.75, ny, 0.045, NH, fill=GOLD)
text(s3, 1.10, ny + 0.15, 5.20, 0.86,
     [B("Also sourced to order", 9.2, True, GREEN_D, HEAD, 4, 1.0),
      B("Cloves (cengkeh), Java long pepper (cabe jawa) and cubeb (kemukus), all whole "
        "dried, from the same farmer network. Volume confirmed against enquiry.",
        8.3, False, GREY, BODY, 0, 1.22)])

rect(s3, 6.86, ny, 5.72, NH, fill=GREEN_L)
rect(s3, 6.86, ny, 0.045, NH, fill=GREEN)
text(s3, 7.21, ny + 0.15, 5.20, 0.86,
     [B("Confirmed per order", 9.2, True, GREEN_D, HEAD, 4, 1.0),
      B("Moisture and foreign matter, packaging and net weight, shelf life, MOQ, incoterms, "
        "port of loading and lead time are agreed against each enquiry. Pricing on application.",
        8.3, False, GREY, BODY, 0, 1.22)])

# --- ajakan + kontak ---
ay = ny + NH + 0.24
AH = 0.94
rect(s3, 0.75, ay, 11.83, AH, fill=GREEN_D)
text(s3, 1.12, ay + 0.15, 7.2, 0.72,
     [B("What we are asking of ATDAG Cairo", 10, True, WHITE, HEAD, 3, 1.0),
      B("Introductions to Egyptian importers, grinders and distributors whose requirements match "
        "this portfolio. We would open any relationship with a trial container, so quality is "
        "judged on the product rather than on paper.", 8.3, False,
        RGBColor(0xD2, 0xDE, 0xD4), BODY, 0, 1.22)])
text(s3, 8.62, ay + 0.16, 3.80, 0.72,
     [B("Arif Rachman  ·  Operational Director", 8.4, True, WHITE, BODY, 2, 1.12),
      B("arif@klumbayanfarm.com  ·  +62 853 7887 0007", 8.0, False,
        RGBColor(0xC9, 0xB8, 0x78), BODY, 2, 1.12),
      B("Jln. Wolter Monginsidi No. 159, Bandar Lampung  ·  www.klumbayanfarm.com",
        7.4, False, RGBColor(0xB8, 0xC6, 0xBA), BODY, 0, 1.16)],
     align=PP_ALIGN.RIGHT)

consortium_strip(s3, 6.42, note=False)
text(s3, 0, 7.20, W, 0.22,
     [B("PT Klumbayan Gold Farm  ·  Product Catalogue  ·  Page 3 of 3",
        7, False, GREY, BODY, 0, 1.0)], align=PP_ALIGN.CENTER)


prs.save(OUT)
print("written:", OUT, os.path.getsize(OUT), "bytes")
